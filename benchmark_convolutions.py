#!/usr/bin/env python3
"""
Performance comparison script for MIOpen vs rocMLIR convolutions.

This script runs all ONNX models in utils/conv_onnx_models_fp32 and 
utils/conv_onnx_models_fp16 with both MIOpen and rocMLIR backends, 
then generates an Excel spreadsheet with the performance comparison results.
"""

import os
import subprocess
import re
import glob
import argparse
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import ColorScaleRule
except ImportError:
    print("openpyxl not found. Installing...")
    subprocess.run(["pip", "install", "openpyxl"], check=True)
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import ColorScaleRule


# Configuration
MIGRAPHX_DRIVER = "/home/umayadav/repo/AMDMIGraphX/build/bin/migraphx-driver"
BASE_DIR = "/home/umayadav/repo/uNetConvs"
UTILS_DIR = os.path.join(BASE_DIR, "utils")
FP32_DIR = os.path.join(UTILS_DIR, "conv_onnx_models_fp32")
FP16_DIR = os.path.join(UTILS_DIR, "conv_onnx_models_fp16")
OUTPUT_DIR = BASE_DIR


def run_benchmark(model_path: str, backend: str) -> dict:
    """
    Run migraphx-driver benchmark for a given model and backend.
    
    Args:
        model_path: Path to the ONNX model
        backend: Either 'miopen' or 'rocmlir'
    
    Returns:
        Dictionary with timing results and verification info
    """
    result = {
        "total_time_ms": None,
        "instruction_found": False,
        "error": None,
        "raw_output": ""
    }
    
    # Set environment variables based on backend
    env = os.environ.copy()
    if backend == "miopen":
        env["MIGRAPHX_DISABLE_MLIR"] = "1"
        expected_instruction = "gpu::convolution"
    else:  # rocmlir
        env["MIGRAPHX_MLIR_USE_SPECIFIC_OPS"] = "convolution,fused"
        env["MIGRAPHX_MLIR_TUNE_EXHAUSTIVE"] = "1"
        expected_instruction = "mlir_convolution"
    
    cmd = [MIGRAPHX_DRIVER, "time", "--exhaustive-tune", model_path]
    
    try:
        print(f"  Running {backend}...", end=" ", flush=True)
        process = subprocess.run(
            cmd,
            env=env,
            capture_output=True,
            text=True,
            timeout=600  # 10 minute timeout per model
        )
        
        output = process.stdout + process.stderr
        result["raw_output"] = output
        
        # Check for expected instruction in output
        if expected_instruction in output:
            result["instruction_found"] = True
        
        # Parse total time from output
        # Looking for pattern like "Total time: X.XXXms" or similar
        time_patterns = [
            r"Total time:\s*([\d.]+)\s*ms",
            r"Total time:\s*([\d.]+)ms",
            r"total time:\s*([\d.]+)\s*ms",
            r"Total:\s*([\d.]+)\s*ms",
        ]
        
        for pattern in time_patterns:
            match = re.search(pattern, output, re.IGNORECASE)
            if match:
                result["total_time_ms"] = float(match.group(1))
                break
        
        if result["total_time_ms"] is not None:
            print(f"{result['total_time_ms']:.3f} ms")
        else:
            # Try to find any timing information
            print("Could not parse time")
            if process.returncode != 0:
                result["error"] = f"Exit code: {process.returncode}"
                
    except subprocess.TimeoutExpired:
        result["error"] = "Timeout (>10min)"
        print("TIMEOUT")
    except Exception as e:
        result["error"] = str(e)
        print(f"ERROR: {e}")
    
    return result


def parse_model_params(model_name: str) -> dict:
    """
    Parse convolution parameters from model filename.
    
    Example: conv_01_n1_ic256_oc256_h8_w8_k3x3.onnx
    """
    params = {
        "batch": None,
        "in_channels": None,
        "out_channels": None,
        "height": None,
        "width": None,
        "kernel": None
    }
    
    # Parse batch size (n)
    match = re.search(r"_n(\d+)_", model_name)
    if match:
        params["batch"] = int(match.group(1))
    
    # Parse input channels (ic)
    match = re.search(r"_ic(\d+)_", model_name)
    if match:
        params["in_channels"] = int(match.group(1))
    
    # Parse output channels (oc)
    match = re.search(r"_oc(\d+)_", model_name)
    if match:
        params["out_channels"] = int(match.group(1))
    
    # Parse height (h)
    match = re.search(r"_h(\d+)_", model_name)
    if match:
        params["height"] = int(match.group(1))
    
    # Parse width (w)
    match = re.search(r"_w(\d+)_", model_name)
    if match:
        params["width"] = int(match.group(1))
    
    # Parse kernel size (k)
    match = re.search(r"_k(\d+x\d+)", model_name)
    if match:
        params["kernel"] = match.group(1)
    
    return params


def create_excel_report(results: list, output_path: str, dtype: str):
    """
    Create an Excel spreadsheet with the benchmark results.
    
    Args:
        results: List of dictionaries with benchmark results
        output_path: Path to save the Excel file
        dtype: Data type (fp32 or fp16)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Performance Comparison"
    
    # Define styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Headers
    headers = [
        "Model Name",
        "Data Type",
        "Batch",
        "In Channels",
        "Out Channels",
        "Height",
        "Width",
        "Kernel",
        "MIOpen (ms)",
        "rocMLIR (ms)",
        "Speedup (MIOpen/rocMLIR)",
        "Winner",
        "MIOpen Valid",
        "rocMLIR Valid"
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Data rows
    miopen_wins = 0
    rocmlir_wins = 0
    ties = 0
    
    for row_idx, result in enumerate(results, 2):
        params = result["params"]
        
        row_data = [
            result["model_name"],
            result.get("dtype", dtype),
            params["batch"],
            params["in_channels"],
            params["out_channels"],
            params["height"],
            params["width"],
            params["kernel"],
            result["miopen_time"],
            result["rocmlir_time"],
            None,  # Speedup - calculated below
            None,  # Winner - calculated below
            "✓" if result["miopen_valid"] else "✗",
            "✓" if result["rocmlir_valid"] else "✗"
        ]
        
        # Calculate speedup and winner
        if result["miopen_time"] and result["rocmlir_time"]:
            speedup = result["miopen_time"] / result["rocmlir_time"]
            row_data[10] = round(speedup, 3)
            
            if speedup > 1.05:
                row_data[11] = "rocMLIR"
                rocmlir_wins += 1
            elif speedup < 0.95:
                row_data[11] = "MIOpen"
                miopen_wins += 1
            else:
                row_data[11] = "Tie"
                ties += 1
        elif result["miopen_time"]:
            row_data[11] = "MIOpen (only)"
            miopen_wins += 1
        elif result["rocmlir_time"]:
            row_data[11] = "rocMLIR (only)"
            rocmlir_wins += 1
        else:
            row_data[11] = "N/A"
        
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Color coding for winner column
            if col == 12:  # Winner column
                if value == "rocMLIR":
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                elif value == "MIOpen":
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                elif value == "Tie":
                    cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    
    # Auto-adjust column widths
    for col in range(1, len(headers) + 1):
        max_length = 0
        column_letter = get_column_letter(col)
        for row in range(1, len(results) + 2):
            cell_value = ws.cell(row=row, column=col).value
            if cell_value:
                max_length = max(max_length, len(str(cell_value)))
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Add summary sheet
    summary_ws = wb.create_sheet("Summary")
    
    summary_data = [
        ["Performance Comparison Summary", ""],
        ["", ""],
        ["Data Type", dtype.upper()],
        ["Total Models Tested", len(results)],
        ["", ""],
        ["rocMLIR Wins (>5% faster)", rocmlir_wins],
        ["MIOpen Wins (>5% faster)", miopen_wins],
        ["Ties (within 5%)", ties],
        ["", ""],
        ["Date Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ]
    
    # Calculate averages for valid results
    valid_results = [r for r in results if r["miopen_time"] and r["rocmlir_time"]]
    if valid_results:
        avg_miopen = sum(r["miopen_time"] for r in valid_results) / len(valid_results)
        avg_rocmlir = sum(r["rocmlir_time"] for r in valid_results) / len(valid_results)
        avg_speedup = avg_miopen / avg_rocmlir if avg_rocmlir > 0 else 0
        
        summary_data.extend([
            ["", ""],
            ["Average MIOpen Time (ms)", round(avg_miopen, 3)],
            ["Average rocMLIR Time (ms)", round(avg_rocmlir, 3)],
            ["Average Speedup (MIOpen/rocMLIR)", round(avg_speedup, 3)],
        ])
    
    for row_idx, row_data in enumerate(summary_data, 1):
        for col_idx, value in enumerate(row_data, 1):
            cell = summary_ws.cell(row=row_idx, column=col_idx, value=value)
            if row_idx == 1:
                cell.font = Font(bold=True, size=14)
            elif col_idx == 1:
                cell.font = Font(bold=True)
    
    summary_ws.column_dimensions['A'].width = 35
    summary_ws.column_dimensions['B'].width = 20
    
    wb.save(output_path)
    print(f"\nExcel report saved to: {output_path}")


def create_combined_excel_report(fp32_results: list, fp16_results: list, output_path: str):
    """
    Create a combined Excel spreadsheet with both fp32 and fp16 results.
    
    Args:
        fp32_results: List of fp32 benchmark results
        fp16_results: List of fp16 benchmark results
        output_path: Path to save the Excel file
    """
    wb = openpyxl.Workbook()
    
    # Define styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Headers
    headers = [
        "Model Name",
        "Data Type",
        "Batch",
        "In Channels",
        "Out Channels",
        "Height",
        "Width",
        "Kernel",
        "MIOpen (ms)",
        "rocMLIR (ms)",
        "Speedup (MIOpen/rocMLIR)",
        "Winner",
        "MIOpen Valid",
        "rocMLIR Valid"
    ]
    
    def add_results_to_sheet(ws, results, dtype):
        """Add results to a worksheet."""
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        miopen_wins = 0
        rocmlir_wins = 0
        ties = 0
        
        for row_idx, result in enumerate(results, 2):
            params = result["params"]
            
            row_data = [
                result["model_name"],
                dtype,
                params["batch"],
                params["in_channels"],
                params["out_channels"],
                params["height"],
                params["width"],
                params["kernel"],
                result["miopen_time"],
                result["rocmlir_time"],
                None,
                None,
                "✓" if result["miopen_valid"] else "✗",
                "✓" if result["rocmlir_valid"] else "✗"
            ]
            
            if result["miopen_time"] and result["rocmlir_time"]:
                speedup = result["miopen_time"] / result["rocmlir_time"]
                row_data[10] = round(speedup, 3)
                
                if speedup > 1.05:
                    row_data[11] = "rocMLIR"
                    rocmlir_wins += 1
                elif speedup < 0.95:
                    row_data[11] = "MIOpen"
                    miopen_wins += 1
                else:
                    row_data[11] = "Tie"
                    ties += 1
            elif result["miopen_time"]:
                row_data[11] = "MIOpen (only)"
                miopen_wins += 1
            elif result["rocmlir_time"]:
                row_data[11] = "rocMLIR (only)"
                rocmlir_wins += 1
            else:
                row_data[11] = "N/A"
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
                if col == 12:
                    if value == "rocMLIR":
                        cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                    elif value == "MIOpen":
                        cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    elif value == "Tie":
                        cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        
        # Auto-adjust column widths
        for col in range(1, len(headers) + 1):
            max_length = 0
            column_letter = get_column_letter(col)
            for row in range(1, len(results) + 2):
                cell_value = ws.cell(row=row, column=col).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        return miopen_wins, rocmlir_wins, ties
    
    # FP32 sheet
    ws_fp32 = wb.active
    ws_fp32.title = "FP32 Results"
    fp32_miopen_wins, fp32_rocmlir_wins, fp32_ties = add_results_to_sheet(ws_fp32, fp32_results, "FP32")
    
    # FP16 sheet
    ws_fp16 = wb.create_sheet("FP16 Results")
    fp16_miopen_wins, fp16_rocmlir_wins, fp16_ties = add_results_to_sheet(ws_fp16, fp16_results, "FP16")
    
    # Summary sheet
    summary_ws = wb.create_sheet("Summary")
    
    summary_data = [
        ["Combined Performance Comparison Summary", ""],
        ["", ""],
        ["Date Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["", ""],
        ["=== FP32 Results ===", ""],
        ["Total Models Tested", len(fp32_results)],
        ["rocMLIR Wins (>5% faster)", fp32_rocmlir_wins],
        ["MIOpen Wins (>5% faster)", fp32_miopen_wins],
        ["Ties (within 5%)", fp32_ties],
    ]
    
    fp32_valid = [r for r in fp32_results if r["miopen_time"] and r["rocmlir_time"]]
    if fp32_valid:
        avg_miopen = sum(r["miopen_time"] for r in fp32_valid) / len(fp32_valid)
        avg_rocmlir = sum(r["rocmlir_time"] for r in fp32_valid) / len(fp32_valid)
        avg_speedup = avg_miopen / avg_rocmlir if avg_rocmlir > 0 else 0
        summary_data.extend([
            ["Average MIOpen Time (ms)", round(avg_miopen, 3)],
            ["Average rocMLIR Time (ms)", round(avg_rocmlir, 3)],
            ["Average Speedup (MIOpen/rocMLIR)", round(avg_speedup, 3)],
        ])
    
    summary_data.extend([
        ["", ""],
        ["=== FP16 Results ===", ""],
        ["Total Models Tested", len(fp16_results)],
        ["rocMLIR Wins (>5% faster)", fp16_rocmlir_wins],
        ["MIOpen Wins (>5% faster)", fp16_miopen_wins],
        ["Ties (within 5%)", fp16_ties],
    ])
    
    fp16_valid = [r for r in fp16_results if r["miopen_time"] and r["rocmlir_time"]]
    if fp16_valid:
        avg_miopen = sum(r["miopen_time"] for r in fp16_valid) / len(fp16_valid)
        avg_rocmlir = sum(r["rocmlir_time"] for r in fp16_valid) / len(fp16_valid)
        avg_speedup = avg_miopen / avg_rocmlir if avg_rocmlir > 0 else 0
        summary_data.extend([
            ["Average MIOpen Time (ms)", round(avg_miopen, 3)],
            ["Average rocMLIR Time (ms)", round(avg_rocmlir, 3)],
            ["Average Speedup (MIOpen/rocMLIR)", round(avg_speedup, 3)],
        ])
    
    for row_idx, row_data in enumerate(summary_data, 1):
        for col_idx, value in enumerate(row_data, 1):
            cell = summary_ws.cell(row=row_idx, column=col_idx, value=value)
            if row_idx == 1 or (isinstance(value, str) and value.startswith("===")):
                cell.font = Font(bold=True, size=14 if row_idx == 1 else 12)
            elif col_idx == 1:
                cell.font = Font(bold=True)
    
    summary_ws.column_dimensions['A'].width = 35
    summary_ws.column_dimensions['B'].width = 20
    
    wb.save(output_path)
    print(f"\nCombined Excel report saved to: {output_path}")


def benchmark_models(models_dir: str, dtype: str) -> list:
    """
    Benchmark all models in a directory.
    
    Args:
        models_dir: Directory containing ONNX models
        dtype: Data type label (fp32 or fp16)
    
    Returns:
        List of benchmark results
    """
    onnx_files = sorted(glob.glob(os.path.join(models_dir, "*.onnx")))
    
    if not onnx_files:
        print(f"No ONNX files found in {models_dir}")
        return []
    
    print(f"\nFound {len(onnx_files)} {dtype.upper()} ONNX models to benchmark")
    print("-" * 70)
    
    results = []
    
    for idx, model_path in enumerate(onnx_files, 1):
        model_name = os.path.basename(model_path)
        print(f"\n[{idx}/{len(onnx_files)}] [{dtype.upper()}] {model_name}")
        
        # Parse model parameters from filename
        params = parse_model_params(model_name)
        
        # Run MIOpen benchmark
        miopen_result = run_benchmark(model_path, "miopen")
        
        # Run rocMLIR benchmark
        rocmlir_result = run_benchmark(model_path, "rocmlir")
        
        results.append({
            "model_name": model_name,
            "dtype": dtype,
            "params": params,
            "miopen_time": miopen_result["total_time_ms"],
            "rocmlir_time": rocmlir_result["total_time_ms"],
            "miopen_valid": miopen_result["instruction_found"],
            "rocmlir_valid": rocmlir_result["instruction_found"],
            "miopen_error": miopen_result["error"],
            "rocmlir_error": rocmlir_result["error"]
        })
    
    return results


def print_summary(results: list, dtype: str):
    """Print summary for a set of results."""
    print(f"\n{dtype.upper()} SUMMARY")
    print("-" * 40)
    
    valid_results = [r for r in results if r["miopen_time"] and r["rocmlir_time"]]
    print(f"Successfully benchmarked: {len(valid_results)}/{len(results)} models")
    
    if valid_results:
        rocmlir_wins = sum(1 for r in valid_results if r["miopen_time"] / r["rocmlir_time"] > 1.05)
        miopen_wins = sum(1 for r in valid_results if r["miopen_time"] / r["rocmlir_time"] < 0.95)
        ties = len(valid_results) - rocmlir_wins - miopen_wins
        
        print(f"rocMLIR wins: {rocmlir_wins}")
        print(f"MIOpen wins: {miopen_wins}")
        print(f"Ties (within 5%): {ties}")
        
        avg_speedup = sum(r["miopen_time"] / r["rocmlir_time"] for r in valid_results) / len(valid_results)
        print(f"Average speedup (MIOpen/rocMLIR): {avg_speedup:.3f}x")


def main():
    parser = argparse.ArgumentParser(
        description='MIOpen vs rocMLIR Convolution Performance Comparison'
    )
    parser.add_argument('--dtype', choices=['fp32', 'fp16', 'both'], default='both',
                        help='Data type to benchmark (default: both)')
    parser.add_argument('--fp32-dir', type=str, default=FP32_DIR,
                        help=f'Directory containing FP32 ONNX models (default: {FP32_DIR})')
    parser.add_argument('--fp16-dir', type=str, default=FP16_DIR,
                        help=f'Directory containing FP16 ONNX models (default: {FP16_DIR})')
    parser.add_argument('--output-dir', type=str, default=OUTPUT_DIR,
                        help=f'Output directory for Excel files (default: {OUTPUT_DIR})')
    
    args = parser.parse_args()
    
    print("=" * 70)
    print("MIOpen vs rocMLIR Convolution Performance Comparison")
    print("=" * 70)
    print(f"migraphx-driver: {MIGRAPHX_DRIVER}")
    
    # Check if migraphx-driver exists
    if not os.path.exists(MIGRAPHX_DRIVER):
        print(f"ERROR: migraphx-driver not found at {MIGRAPHX_DRIVER}")
        return
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    fp32_results = []
    fp16_results = []
    
    # Benchmark FP32 models
    if args.dtype in ['fp32', 'both']:
        print("\n" + "=" * 70)
        print("FP32 BENCHMARKS")
        print("=" * 70)
        fp32_results = benchmark_models(args.fp32_dir, "fp32")
        
        if fp32_results:
            # Generate FP32-only report
            fp32_output = os.path.join(args.output_dir, f"convolution_benchmark_fp32_{timestamp}.xlsx")
            create_excel_report(fp32_results, fp32_output, "fp32")
            print_summary(fp32_results, "fp32")
    
    # Benchmark FP16 models
    if args.dtype in ['fp16', 'both']:
        print("\n" + "=" * 70)
        print("FP16 BENCHMARKS")
        print("=" * 70)
        fp16_results = benchmark_models(args.fp16_dir, "fp16")
        
        if fp16_results:
            # Generate FP16-only report
            fp16_output = os.path.join(args.output_dir, f"convolution_benchmark_fp16_{timestamp}.xlsx")
            create_excel_report(fp16_results, fp16_output, "fp16")
            print_summary(fp16_results, "fp16")
    
    # Generate combined report if both were run
    if fp32_results and fp16_results:
        combined_output = os.path.join(args.output_dir, f"convolution_benchmark_combined_{timestamp}.xlsx")
        create_combined_excel_report(fp32_results, fp16_results, combined_output)
    
    # Print final summary
    print("\n" + "=" * 70)
    print("FINAL SUMMARY")
    print("=" * 70)
    
    if fp32_results:
        print_summary(fp32_results, "fp32")
    
    if fp16_results:
        print_summary(fp16_results, "fp16")
    
    print("\n✓ Done!")


if __name__ == "__main__":
    main()
